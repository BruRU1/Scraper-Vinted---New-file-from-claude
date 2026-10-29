import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

SAMPLE = len(sys.argv) > 1 and sys.argv[1] == "sample"
OUT = "UK_Sole_Trader_Tax_Tracker_SAMPLE.xlsx" if SAMPLE else "UK_Sole_Trader_Tax_Tracker.xlsx"
N = 1000  # rows of capacity
NAVY, TEAL, LIGHT, INPUT = "1F3A5F", "2A9D8F", "EAF4F2", "FFF8E1"
hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor=NAVY)
inp_fill = PatternFill("solid", fgColor=INPUT); tot_fill = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="CCCCCC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
GBP = '£#,##0.00;[Red]-£#,##0.00'

wb = Workbook()

def header(ws, row, labels, widths=None):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l); c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    if sub: ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="666666")

# ---------- Start Here ----------
s = wb.active; s.title = "Start Here"
title(s, "UK Sole Trader Tax Tracker", "Income & expenses log, quarterly (MTD) summaries and a tax estimate - Excel & Google Sheets")
lines = [
 ("How to use", True),
 ("1. Go to Settings and check the tax year start date (default 6 April 2026 - 5 April 2027).", False),
 ("2. Log every payment you receive on the Income sheet (yellow cells only).", False),
 ("3. Log every business cost on the Expenses sheet and pick a category from the drop-down.", False),
 ("4. The Quarterly Summary sheet totals income and expenses per quarter - use these figures for your MTD quarterly updates.", False),
 ("5. The Tax Estimate sheet shows an approximate Income Tax + Class 4 NIC bill for the year.", False),
 ("", False),
 ("Important notes", True),
 ("- Quarters follow the standard tax-year periods: 6 Apr-5 Jul, 6 Jul-5 Oct, 6 Oct-5 Jan, 6 Jan-5 Apr.", False),
 ("- The estimate assumes your only income is this self-employment, and uses England/Wales/Northern Ireland rates. Scottish rates differ.", False),
 ("- It is an ESTIMATE to help you plan, not tax advice. Check figures against GOV.UK or an accountant before filing.", False),
 ("- This workbook does not submit anything to HMRC. For MTD you also need HMRC-recognised software (many accept spreadsheet uploads via bridging software).", False),
 ("- Rates and thresholds live on the Settings sheet - update them if the rules change.", False),
 ("- Yellow cells = you type here. Everything else is calculated; please don't overwrite it.", False),
]
for i, (t, b) in enumerate(lines, 4):
    s.cell(row=i, column=1, value=t).font = Font(bold=b, size=12 if b else 11, color=NAVY if b else "000000")
s.column_dimensions["A"].width = 130

# ---------- Settings ----------
st = wb.create_sheet("Settings")
title(st, "Settings", "Change the yellow cells only if the tax year or rules change")
rows = [
 ("Tax year start", dt.date(2026, 4, 6), "TaxStart", "dd/mm/yyyy"),
 ("Tax year end", "=DATE(YEAR(TaxStart)+1,4,5)", "TaxEnd", "dd/mm/yyyy"),
 ("Personal allowance", 12570, "PA", GBP),
 ("Allowance taper starts at income of", 100000, "TaperStart", GBP),
 ("Basic rate band (taxable income)", 37700, "BasicBand", GBP),
 ("Additional rate threshold (taxable income)", 125140, "AddlThreshold", GBP),
 ("Basic rate", 0.20, "BasicRate", "0%"),
 ("Higher rate", 0.40, "HigherRate", "0%"),
 ("Additional rate", 0.45, "AddlRate", "0%"),
 ("Class 4 NIC lower profits limit", 12570, "NICLower", GBP),
 ("Class 4 NIC upper profits limit", 50270, "NICUpper", GBP),
 ("Class 4 NIC main rate", 0.06, "NICMain", "0%"),
 ("Class 4 NIC rate above upper limit", 0.02, "NICUpperRate", "0%"),
 ("Trading allowance (alternative to claiming expenses)", 1000, "TradingAllow", GBP),
]
header(st, 4, ["Setting", "Value"], [52, 18])
for i, (lab, val, name, fmt) in enumerate(rows, 5):
    st.cell(row=i, column=1, value=lab).border = box
    c = st.cell(row=i, column=2, value=val); c.number_format = fmt; c.border = box
    if not (isinstance(val, str) and val.startswith("=")): c.fill = inp_fill
    wb.defined_names[name] = DefinedName(name, attr_text=f"Settings!$B${i}")

# ---------- Income ----------
inc = wb.create_sheet("Income")
title(inc, "Income", "Money received from customers (before any expenses)")
header(inc, 4, ["Date", "Customer", "Description", "Amount (£)", "Quarter"], [14, 28, 40, 16, 18])
for r in range(5, 5 + N):
    for c in (1, 2, 3, 4):
        cell = inc.cell(row=r, column=c); cell.fill = inp_fill; cell.border = box
    inc.cell(row=r, column=1).number_format = "dd/mm/yyyy"; inc.cell(row=r, column=4).number_format = GBP
    q = inc.cell(row=r, column=5)
    q.value = (f'=IF(A{r}="","",IF(OR(A{r}<TaxStart,A{r}>TaxEnd),"Outside tax year",'
               f'IF(A{r}<DATE(YEAR(TaxStart),7,6),"Q1",IF(A{r}<DATE(YEAR(TaxStart),10,6),"Q2",'
               f'IF(A{r}<DATE(YEAR(TaxStart)+1,1,6),"Q3","Q4")))))')
    q.border = box
inc.freeze_panes = "A5"

# ---------- Expenses ----------
CATS = ["Cost of goods bought for resale", "Subcontractors / staff costs", "Car, van & travel",
        "Premises & utilities", "Office, phone & software", "Advertising & marketing",
        "Bank & finance charges", "Professional fees (accountant etc.)", "Training & subscriptions",
        "Insurance", "Other allowable expenses"]
exp = wb.create_sheet("Expenses")
title(exp, "Expenses", "Allowable business costs - pick a category from the drop-down")
header(exp, 4, ["Date", "Supplier", "Description", "Category", "Amount (£)", "Quarter"], [14, 28, 40, 36, 16, 18])
dv = DataValidation(type="list", formula1="=Settings!$D$5:$D$15", allow_blank=True); exp.add_data_validation(dv)
for r in range(5, 5 + N):
    for c in (1, 2, 3, 4, 5):
        cell = exp.cell(row=r, column=c); cell.fill = inp_fill; cell.border = box
    exp.cell(row=r, column=1).number_format = "dd/mm/yyyy"; exp.cell(row=r, column=5).number_format = GBP
    dv.add(exp.cell(row=r, column=4))
    q = exp.cell(row=r, column=6)
    q.value = (f'=IF(A{r}="","",IF(OR(A{r}<TaxStart,A{r}>TaxEnd),"Outside tax year",'
               f'IF(A{r}<DATE(YEAR(TaxStart),7,6),"Q1",IF(A{r}<DATE(YEAR(TaxStart),10,6),"Q2",'
               f'IF(A{r}<DATE(YEAR(TaxStart)+1,1,6),"Q3","Q4")))))')
    q.border = box
exp.freeze_panes = "A5"
st["D4"] = "Expense categories"; st["D4"].font = hdr_font; st["D4"].fill = hdr_fill; st.column_dimensions["D"].width = 38
for i, c in enumerate(CATS, 5): st.cell(row=i, column=4, value=c).fill = inp_fill

# ---------- Quarterly Summary ----------
qs = wb.create_sheet("Quarterly Summary")
title(qs, "Quarterly Summary (MTD)", "Use these totals for your quarterly updates")
header(qs, 4, ["", "Q1", "Q2", "Q3", "Q4", "Full year"], [40, 16, 16, 16, 16, 18])
qs["A5"] = "Period"
qs["B5"] = '=TEXT(TaxStart,"d mmm yy")&" - "&TEXT(DATE(YEAR(TaxStart),7,5),"d mmm yy")'
qs["C5"] = '=TEXT(DATE(YEAR(TaxStart),7,6),"d mmm yy")&" - "&TEXT(DATE(YEAR(TaxStart),10,5),"d mmm yy")'
qs["D5"] = '=TEXT(DATE(YEAR(TaxStart),10,6),"d mmm yy")&" - "&TEXT(DATE(YEAR(TaxStart)+1,1,5),"d mmm yy")'
qs["E5"] = '=TEXT(DATE(YEAR(TaxStart)+1,1,6),"d mmm yy")&" - "&TEXT(TaxEnd,"d mmm yy")'
qs["F5"] = '=TEXT(TaxStart,"d mmm yy")&" - "&TEXT(TaxEnd,"d mmm yy")'
for col in "BCDEF": qs[f"{col}5"].alignment = Alignment(horizontal="center"); qs[f"{col}5"].font = Font(size=9, color="666666")
qs["A6"] = "Total income"; qs["A6"].font = Font(bold=True)
for col in "BCDE":
    qs[f"{col}6"] = f'=SUMIFS(Income!$D$5:$D${4+N},Income!$E$5:$E${4+N},{col}$4)'
qs["F6"] = "=SUM(B6:E6)"
qs["A8"] = "Expenses by category"; qs["A8"].font = Font(bold=True, color=NAVY)
for i, c in enumerate(CATS):
    r = 9 + i
    qs.cell(row=r, column=1, value=f"=Settings!D{5+i}")
    for col in "BCDE":
        qs[f"{col}{r}"] = f'=SUMIFS(Expenses!$E$5:$E${4+N},Expenses!$D$5:$D${4+N},$A{r},Expenses!$F$5:$F${4+N},{col}$4)'
    qs[f"F{r}"] = f"=SUM(B{r}:E{r})"
last = 9 + len(CATS) - 1; tr = last + 1
qs[f"A{tr}"] = "Total expenses"; qs[f"A{tr}"].font = Font(bold=True)
for col in "BCDEF": qs[f"{col}{tr}"] = f"=SUM({col}9:{col}{last})"
pr = tr + 2
qs[f"A{pr}"] = "Net profit (income - expenses)"; qs[f"A{pr}"].font = Font(bold=True)
for col in "BCDEF": qs[f"{col}{pr}"] = f"={col}6-{col}{tr}"
qs[f"A{pr+2}"] = "Check: entries dated outside the tax year (should be 0)"
qs[f"F{pr+2}"] = (f'=COUNTIF(Income!$E$5:$E${4+N},"Outside tax year")+COUNTIF(Expenses!$F$5:$F${4+N},"Outside tax year")')
for r in range(6, pr + 1):
    for col in "BCDEF":
        qs[f"{col}{r}"].number_format = GBP
        if r in (6, tr, pr): qs[f"{col}{r}"].fill = tot_fill; qs[f"{col}{r}"].font = Font(bold=True)
qs.freeze_panes = "B5"
SUMROW_INC, SUMROW_EXP = 6, tr

# ---------- Tax Estimate ----------
te = wb.create_sheet("Tax Estimate")
title(te, "Tax Estimate", "Approximate Income Tax + Class 4 NIC on self-employment profit (England, Wales & NI)")
te.column_dimensions["A"].width = 58; te.column_dimensions["B"].width = 18; te.column_dimensions["C"].width = 60
items = [
 ("Total income", f"='Quarterly Summary'!F{SUMROW_INC}", ""),
 ("Total expenses claimed", f"='Quarterly Summary'!F{SUMROW_EXP}", ""),
 ("Deduction used", "=MAX(B6,MIN(B5,TradingAllow))", "Higher of your expenses or the trading allowance (max £1,000)"),
 ("Taxable profit", "=MAX(0,B5-B7)", ""),
 ("Personal allowance (after taper)", "=MAX(0,PA-MAX(0,(B8-TaperStart)/2))", "Reduces £1 for every £2 over £100,000"),
 ("Taxable income after allowance", "=MAX(0,B8-B9)", ""),
 ("Tax at basic rate", "=BasicRate*MIN(B10,BasicBand)", ""),
 ("Tax at higher rate", "=HigherRate*MAX(0,MIN(B10,AddlThreshold)-BasicBand)", ""),
 ("Tax at additional rate", "=AddlRate*MAX(0,B10-AddlThreshold)", ""),
 ("Income Tax", "=SUM(B11:B13)", ""),
 ("Class 4 NIC", "=IF(B8<=NICLower,0,NICMain*(MIN(B8,NICUpper)-NICLower)+NICUpperRate*MAX(0,B8-NICUpper))", "Class 2 is no longer compulsory"),
 ("Total estimated tax + NIC", "=B14+B15", ""),
 ("Suggested amount to set aside each month", "=B16/12", ""),
 ("Effective rate on profit", "=IF(B8=0,0,B16/B8)", ""),
]
header(te, 4, ["", "Amount", "Notes"])
for i, (lab, f, note) in enumerate(items, 5):
    te.cell(row=i, column=1, value=lab); c = te.cell(row=i, column=2, value=f)
    c.number_format = "0.0%" if lab.startswith("Effective") else GBP
    te.cell(row=i, column=3, value=note).font = Font(italic=True, color="666666")
for r in (8, 14, 16):
    te.cell(row=r, column=1).font = Font(bold=True); te.cell(row=r, column=2).font = Font(bold=True); te.cell(row=r, column=2).fill = tot_fill
te["A20"] = "Estimate only - not tax advice. Check GOV.UK or ask an accountant before filing."; te["A20"].font = Font(italic=True, color="B00020")

# ---------- sample data ----------
if SAMPLE:
    D = dt.date
    for i, (d, c, desc, a) in enumerate([
        (D(2026,4,20),"Acme Ltd","Website design",1500),(D(2026,6,30),"Bloom Cafe","Logo",800),
        (D(2026,7,6),"Acme Ltd","Maintenance",600),(D(2026,10,5),"J Smith","Consulting",2500),
        (D(2026,10,6),"Bloom Cafe","Menu design",1200),(D(2027,1,5),"Acme Ltd","Retainer",3000),
        (D(2027,1,6),"Acme Ltd","Retainer",3000),(D(2027,3,31),"J Smith","Project",8000),
        (D(2027,4,6),"Late Co","Next year job",999)], 5):
        inc.cell(row=i,column=1,value=d); inc.cell(row=i,column=2,value=c); inc.cell(row=i,column=3,value=desc); inc.cell(row=i,column=4,value=a)
    for i, (d, s_, desc, cat, a) in enumerate([
        (D(2026,5,1),"Adobe","Subscription",CATS[4],200),(D(2026,8,1),"Google","Ads",CATS[5],300),
        (D(2026,11,1),"Train Co","Travel",CATS[2],150),(D(2027,2,1),"Accountant","Fees",CATS[7],350)], 5):
        exp.cell(row=i,column=1,value=d); exp.cell(row=i,column=2,value=s_); exp.cell(row=i,column=3,value=desc)
        exp.cell(row=i,column=4,value=cat); exp.cell(row=i,column=5,value=a)

wb.save(OUT); print("saved", OUT)
