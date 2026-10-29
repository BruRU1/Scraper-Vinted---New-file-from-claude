import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
SAMPLE = len(sys.argv) > 1 and sys.argv[1] == "sample"
OUT = "UK_Freelancer_Invoice_Pack_SAMPLE.xlsx" if SAMPLE else "UK_Freelancer_Invoice_Pack.xlsx"
NAVY, INPUT, LIGHT = "1F3A5F", "FFF8E1", "EAF4F2"
hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor=NAVY)
inp = PatternFill("solid", fgColor=INPUT); tot = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="CCCCCC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
GBP = '£#,##0.00;[Red]-£#,##0.00'
wb = Workbook()

# Settings
st = wb.active; st.title = "Settings"
st["A1"] = "Your business details"; st["A1"].font = Font(bold=True, size=16, color=NAVY)
st["A2"] = "Fill the yellow cells once - they appear on every invoice"; st["A2"].font = Font(italic=True, color="666666")
rows = [("Business / trading name", "[YOUR BUSINESS NAME]", "BizName"), ("Address line 1", "[ADDRESS LINE 1]", "Addr1"),
 ("Address line 2 / postcode", "[TOWN, POSTCODE]", "Addr2"), ("Email", "[YOUR EMAIL]", "BizEmail"),
 ("Phone", "[YOUR PHONE]", "BizPhone"), ("VAT number (leave blank if not VAT registered)", "", "VATNo"),
 ("VAT rate (0% if not registered)", 0, "VATRate"), ("Bank account name", "[ACCOUNT NAME]", "BankName"),
 ("Sort code", "[00-00-00]", "SortCode"), ("Account number", "[00000000]", "AcctNo"),
 ("Payment terms (days)", 14, "TermsDays"), ("Next invoice number", 1001, "NextNo")]
for i, (l, v, n) in enumerate(rows, 4):
    st.cell(row=i, column=1, value=l).border = box
    c = st.cell(row=i, column=2, value=v); c.fill = inp; c.border = box
    if n == "VATRate": c.number_format = "0%"
    wb.defined_names[n] = DefinedName(n, attr_text=f"Settings!$B${i}")
st.column_dimensions["A"].width = 52; st.column_dimensions["B"].width = 32

# Invoice
iv = wb.create_sheet("Invoice")
for col, w in zip("ABCDE", (46, 10, 16, 16, 4)): iv.column_dimensions[col].width = w
iv["A1"] = "=BizName"; iv["A1"].font = Font(bold=True, size=18, color=NAVY)
iv["A2"] = "=Addr1"; iv["A3"] = "=Addr2"; iv["A4"] = '=BizEmail&"  |  "&BizPhone'
iv["A5"] = '=IF(VATNo="","","VAT No: "&VATNo)'
iv["C1"] = "INVOICE"; iv["C1"].font = Font(bold=True, size=18, color=NAVY)
iv["C3"] = "Invoice no."; iv["D3"] = "=NextNo"; iv["D3"].fill = inp
iv["C4"] = "Invoice date"; iv["D4"] = dt.date(2026, 11, 2) if SAMPLE else None; iv["D4"].fill = inp; iv["D4"].number_format = "dd/mm/yyyy"
iv["C5"] = "Due date"; iv["D5"] = '=IF(D4="","",D4+TermsDays)'; iv["D5"].number_format = "dd/mm/yyyy"
iv["A7"] = "Bill to"; iv["A7"].font = Font(bold=True)
iv["A8"] = "[CLIENT NAME]" if not SAMPLE else "Acme Ltd"; iv["A9"] = "[CLIENT ADDRESS]" if not SAMPLE else "1 High Street, London"
for a in ("A8", "A9"): iv[a].fill = inp
for i, h in enumerate(["Description", "Qty", "Unit price", "Amount"], 1):
    c = iv.cell(row=11, column=i, value=h); c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center")
for r in range(12, 22):
    for c in range(1, 4): iv.cell(row=r, column=c).fill = inp; iv.cell(row=r, column=c).border = box
    iv.cell(row=r, column=3).number_format = GBP
    iv.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",B{r}*C{r})').number_format = GBP
    iv.cell(row=r, column=4).border = box
iv["C23"] = "Subtotal"; iv["D23"] = "=SUM(D12:D21)"
iv["C24"] = '="VAT @ "&TEXT(VATRate,"0%")'; iv["D24"] = "=ROUND(D23*VATRate,2)"
iv["C25"] = "TOTAL DUE"; iv["D25"] = "=D23+D24"
for r in (23, 24, 25): iv[f"D{r}"].number_format = GBP
for a in ("C25", "D25"): iv[a].font = Font(bold=True, size=12); iv[a].fill = tot
iv["A23"] = '="Payment due within "&TermsDays&" days"'
iv["A25"] = '="Pay by bank transfer to "&BankName&", sort code "&SortCode&", account "&AcctNo&", ref INV-"&D3'
iv["A25"].alignment = Alignment(wrap_text=True, vertical="top")
iv.merge_cells("A25:B27"); iv.print_area = "A1:D28"; iv.page_setup.fitToWidth = 1
if SAMPLE:
    for r, (d, q, p) in enumerate([("Website design", 1, 1500), ("Hosting (12 months)", 12, 10)], 12):
        iv.cell(row=r, column=1, value=d); iv.cell(row=r, column=2, value=q); iv.cell(row=r, column=3, value=p)

# Log
lg = wb.create_sheet("Invoice Log")
lg["A1"] = "Invoice Log"; lg["A1"].font = Font(bold=True, size=16, color=NAVY)
lg["A2"] = "Add a row for each invoice sent. Status turns OVERDUE automatically."; lg["A2"].font = Font(italic=True, color="666666")
heads = ["Invoice no.", "Client", "Issue date", "Terms (days)", "Due date", "Amount (£)", "Date paid", "Status", "Days overdue"]
for i, h in enumerate(heads, 1):
    c = lg.cell(row=4, column=i, value=h); c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center", wrap_text=True)
for col, w in zip("ABCDEFGHI", (12, 28, 14, 12, 14, 14, 14, 14, 14)): lg.column_dimensions[col].width = w
for r in range(5, 505):
    for c in (1, 2, 3, 4, 6, 7): lg.cell(row=r, column=c).fill = inp; lg.cell(row=r, column=c).border = box
    for c in (3, 5, 7): lg.cell(row=r, column=c).number_format = "dd/mm/yyyy"
    lg.cell(row=r, column=6).number_format = GBP
    lg.cell(row=r, column=5, value=f'=IF(OR(C{r}="",D{r}=""),"",C{r}+D{r})')
    lg.cell(row=r, column=8, value=f'=IF(A{r}="","",IF(G{r}<>"","Paid",IF(TODAY()>E{r},"OVERDUE","Awaiting payment")))')
    lg.cell(row=r, column=9, value=f'=IF(H{r}="OVERDUE",TODAY()-E{r},"")')
lg["K4"] = "Outstanding"; lg["K5"] = '=SUMIFS(F5:F504,H5:H504,"Awaiting payment")+SUMIFS(F5:F504,H5:H504,"OVERDUE")'
lg["K7"] = "Overdue"; lg["K8"] = '=SUMIFS(F5:F504,H5:H504,"OVERDUE")'
lg["K10"] = "Paid"; lg["K11"] = '=SUMIFS(F5:F504,H5:H504,"Paid")'
for a in ("K4", "K7", "K10"): lg[a].font = Font(bold=True)
for a in ("K5", "K8", "K11"): lg[a].number_format = GBP; lg[a].fill = tot
lg.column_dimensions["K"].width = 16; lg.freeze_panes = "A5"
from openpyxl.formatting.rule import FormulaRule
lg.conditional_formatting.add("A5:I504", FormulaRule(formula=['$H5="OVERDUE"'], fill=PatternFill("solid", bgColor="F8D7DA")))
lg.conditional_formatting.add("A5:I504", FormulaRule(formula=['$H5="Paid"'], fill=PatternFill("solid", bgColor="D4EDDA")))
if SAMPLE:
    D = dt.date
    for r, v in enumerate([(1001, "Acme Ltd", D(2026,9,1), 14, None, 1620, D(2026,9,10)), (1002, "Bloom Cafe", D(2026,9,20), 14, None, 800, None),
                           (1003, "J Smith", D(2099,1,1), 14, None, 500, None)], 5):
        for c, x in enumerate(v, 1):
            if x is not None: lg.cell(row=r, column=c, value=x)
wb.save(OUT); print("saved", OUT)
