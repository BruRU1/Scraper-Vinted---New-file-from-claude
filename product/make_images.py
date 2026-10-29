import datetime as dt, html
from openpyxl import load_workbook
import os
os.makedirs('_html',exist_ok=True)
NAVY,TEAL="#1F3A5F","#2A9D8F"
def fmt(v,money=False):
    if v is None or v=="": return ""
    if isinstance(v,(dt.datetime,dt.date)): return v.strftime("%d/%m/%Y")
    if isinstance(v,(int,float)): return f"£{v:,.2f}" if money else f"{v:g}"
    return html.escape(str(v))
def table(ws,r1,r2,c2,money_cols=(),bold_rows=(),head=True):
    out="<table>"
    for r in range(r1,r2+1):
        cells=[ws.cell(row=r,column=c).value for c in range(1,c2+1)]
        if all(x in (None,"") for x in cells): continue
        tag="th" if (head and r==r1) else "td"
        cls=' class="b"' if r in bold_rows else ""
        out+=f"<tr{cls}>"+"".join(f"<{tag}>{fmt(v,(c+1 in money_cols) and tag=='td')}</{tag}>" for c,v in enumerate(cells))+"</tr>"
    return out+"</table>"
t=load_workbook("_tmp/UK_Sole_Trader_Tax_Tracker_SAMPLE.xlsx",data_only=True)
i=load_workbook("_tmp/UK_Freelancer_Invoice_Pack_SAMPLE.xlsx",data_only=True)
qs=t["Quarterly Summary"]; te=t["Tax Estimate"]; inv=i["Invoice"]; lg=i["Invoice Log"]
css=f"""*{{box-sizing:border-box;font-family:Helvetica,Arial,sans-serif}}body{{margin:0;width:1600px;height:1600px;background:#fff;color:#222}}
.hero{{display:flex;flex-direction:column;justify-content:center;height:1600px;background:linear-gradient(160deg,{NAVY},{TEAL});color:#fff;padding:110px 100px;position:relative}}
.hero h1{{font-size:128px;line-height:1.05;margin:0 0 30px}}.hero h2{{font-size:46px;font-weight:400;margin:0 0 60px;opacity:.95}}
.hero li{{font-size:54px;margin:26px 0;list-style:none}}.hero li:before{{content:"✓ ";color:#FFD166}}
.badge{{position:absolute;right:90px;bottom:90px;background:#FFD166;color:{NAVY};font-weight:bold;font-size:44px;padding:28px 44px;border-radius:20px}}
.slide{{padding:90px;height:1600px;background:#f6f8fa}}.slide h1{{color:{NAVY};font-size:74px;margin:0 0 14px}}.slide p{{font-size:38px;color:#555;margin:0 0 50px}}
table{{border-collapse:collapse;width:100%;background:#fff;box-shadow:0 8px 30px rgba(0,0,0,.12);font-size:30px}}
th{{background:{NAVY};color:#fff;padding:16px;text-align:center}}td{{padding:13px 18px;border:1px solid #e3e3e3;text-align:right}}td:first-child{{text-align:left}}
tr.b td{{background:#EAF4F2;font-weight:bold}}.note{{margin-top:40px;font-size:32px;color:{NAVY}}}
.inv{{background:#fff;padding:60px;box-shadow:0 8px 30px rgba(0,0,0,.12)}}.inv h2{{margin:0;color:{NAVY};font-size:52px}}
"""
def page(body): return f"<html><head><style>{css}</style></head><body>{body}</body></html>"
def hero(title,sub,bul,badge): return page(f'<div class="hero"><h1>{title}</h1><h2>{sub}</h2><ul style="padding:0">'+"".join(f"<li>{b}</li>" for b in bul)+f'</ul><div class="badge">{badge}</div></div>')
pages={}
pages["tracker_1_cover"]=hero("UK Sole Trader<br>Tax Tracker","Income &amp; expenses, MTD quarterly summaries and tax estimate",["Excel + Google Sheets","Quarterly (MTD) totals in seconds","Income Tax + Class 4 NIC estimate","No macros, no subscription"],"2026/27")
pages["tracker_2_quarterly"]=page(f'<div class="slide"><h1>Quarterly Summary</h1><p>Income and expenses per MTD quarter, split by category - automatically.</p>{table(qs,4,qs.max_row-3,6,money_cols=range(2,7),bold_rows=(6,20,22))}<div class="note">Entries dated outside the tax year are flagged for you.</div></div>')
pages["tracker_3_tax"]=page(f'<div class="slide"><h1>Know your tax bill</h1><p>Estimated Income Tax + Class 4 NIC, with a monthly "set aside" figure.</p>{table(te,4,18,2,money_cols=(2,),bold_rows=(8,14,16))}<div class="note">Estimate only - not tax advice. England, Wales &amp; NI rates.</div></div>')
pages["invoice_1_cover"]=hero("UK Freelancer<br>Invoice Pack","Professional invoices + an overdue payment tracker",["Auto-totals, optional VAT","Your details entered once","Overdue invoices flagged automatically","Excel + Google Sheets"],"Printable")
rows="".join(f"<tr><td>{fmt(inv.cell(row=r,column=1).value)}</td><td>{fmt(inv.cell(row=r,column=2).value)}</td><td>{fmt(inv.cell(row=r,column=3).value,True)}</td><td>{fmt(inv.cell(row=r,column=4).value,True)}</td></tr>" for r in (12,13))
pages["invoice_2_invoice"]=page(f'<div class="slide"><h1>Clean, printable invoices</h1><p>Type the line items - totals, due date and bank details fill themselves in.</p><div class="inv"><div style="display:flex;justify-content:space-between"><div><h2>[YOUR BUSINESS NAME]</h2><div style="font-size:28px;color:#666">[ADDRESS] | [EMAIL]</div></div><div style="text-align:right;font-size:30px"><h2>INVOICE</h2>No. {fmt(inv["D3"].value)}<br>Date {fmt(inv["D4"].value)}<br>Due {fmt(inv["D5"].value)}</div></div><p style="margin-top:50px;font-size:30px"><b>Bill to:</b> Acme Ltd, 1 High Street, London</p><table><tr><th>Description</th><th>Qty</th><th>Unit price</th><th>Amount</th></tr>{rows}<tr class="b"><td></td><td></td><td>TOTAL DUE</td><td>{fmt(inv["D25"].value,True)}</td></tr></table><p style="font-size:26px;margin-top:40px">{html.escape(str(inv["A25"].value))}</p></div></div>')
pages["invoice_3_log"]=page(f'<div class="slide"><h1>Never chase blind again</h1><p>Overdue invoices turn up automatically, with days overdue and totals.</p>{table(lg,4,7,9,money_cols=(6,))}<div class="note">Outstanding {fmt(lg["K5"].value,True)} &middot; Overdue {fmt(lg["K8"].value,True)} &middot; Paid {fmt(lg["K11"].value,True)}</div></div>')
pages["bundle_cover"]=hero("Self-Employed<br>Starter Bundle","Tax tracker + invoice pack - everything to run your books",["UK Sole Trader Tax Tracker","UK Freelancer Invoice Pack","Excel + Google Sheets","Save vs buying separately"],"BUNDLE")
for n,h in pages.items(): open(f'_html/{n}.html','w').write(h)
print(len(pages))
